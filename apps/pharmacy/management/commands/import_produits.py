from datetime import datetime
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.finance.models import get_current_rate
from apps.pharmacy.models import PharmacyProduct, StockMovement


def parse_montant(valeur):
    """Parse '9 200,00 FC' ou '$4,00' ou 9200.0 -> Decimal. Renvoie None si vide/0."""
    if valeur is None:
        return None
    if isinstance(valeur, (int, float, Decimal)):
        d = Decimal(str(valeur))
        return d if d > 0 else None
    txt = str(valeur).strip()
    if not txt:
        return None
    txt = txt.replace('FC', '').replace('$', '').replace('\u00a0', ' ').replace(' ', '').replace(',', '.')
    try:
        d = Decimal(txt)
        return d if d > 0 else None
    except InvalidOperation:
        return None


class Command(BaseCommand):
    help = "Importe les produits pharmaceutiques depuis le fichier Excel du CRF-MK."

    def add_arguments(self, parser):
        parser.add_argument('fichier', type=str, help="Chemin du fichier .xlsx")
        parser.add_argument('--feuille', type=str, default=None, help="Nom de la feuille (optionnel)")

    def handle(self, *args, **options):
        wb = load_workbook(options['fichier'], data_only=True)
        ws = wb[options['feuille']] if options['feuille'] else wb.active
        rate = get_current_rate()
        admin = User.objects.filter(is_staff=True).first()

        crees, existants, ignores, erreurs = 0, 0, 0, 0

        # Colonnes : 0=N° 1=Produit 2=Dosage 3=Forme 4=Quantité 5=Sortie(Usé)
        #            6=Reste 7=PRIX FC 8=PRIX $ 9=Observation
        for ligne in ws.iter_rows(min_row=2, values_only=True):
            if not ligne or not ligne[1]:
                ignores += 1
                continue

            nom_base = str(ligne[1]).strip()
            dosage = str(ligne[2]).strip() if ligne[2] else ''
            forme = str(ligne[3]).strip() if ligne[3] else ''
            nom = f"{nom_base} {dosage}".strip() if dosage else nom_base

            try:
                quantite = int(float(str(ligne[4]).replace(',', '.'))) if ligne[4] else 0
                sortie = int(float(str(ligne[5]).replace(',', '.'))) if ligne[5] else 0
            except (ValueError, TypeError):
                self.stdout.write(self.style.ERROR(f"Quantité invalide : {nom}"))
                erreurs += 1
                continue

            prix_fc = parse_montant(ligne[7])
            prix_usd = parse_montant(ligne[8])

            # Dérivation si une seule devise renseignée
            if prix_fc is None and prix_usd is not None:
                prix_fc = (prix_usd * rate).quantize(Decimal('1'))
            if prix_usd is None and prix_fc is not None:
                prix_usd = (prix_fc / rate).quantize(Decimal('0.01'))

            observation = str(ligne[9]).strip() if len(ligne) > 9 and ligne[9] else ''
            if prix_usd is None:
                observation = (observation + ' · ' if observation else '') + '⚠ Prix à définir'
                prix_usd, prix_fc = Decimal('0'), Decimal('0')
            if forme:
                observation = f"Forme: {forme}" + (f" · {observation}" if observation else '')

            produit, created = PharmacyProduct.objects.get_or_create(
                name__iexact=nom,
                defaults={
                    'name': nom,
                    'price_usd': prix_usd,
                    'price_fc': prix_fc,
                    'observation': observation,
                },
            )
            if not created:
                existants += 1
                continue

            crees += 1
            self.stdout.write(f"  + {nom} — {prix_usd}$ / {prix_fc} FC (stock restant: {quantite - sortie})")

            # Historique tracé : entrée du stock initial + sortie "déjà utilisé"
            if quantite > 0:
                StockMovement.objects.create(
                    product=produit, movement_type=StockMovement.Type.IN,
                    quantity=quantite, reason="Stock initial (import Excel)", created_by=admin)
            if sortie > 0:
                StockMovement.objects.create(
                    product=produit, movement_type=StockMovement.Type.OUT,
                    quantity=sortie, reason="Historique — déjà utilisé (import Excel)",
                    created_by=admin)

        self.stdout.write(self.style.SUCCESS(
            f"\nTerminé : {crees} produits créés, {existants} déjà existants, "
            f"{ignores} lignes vides ignorées, {erreurs} erreurs. Taux FC : {rate}."
        ))s